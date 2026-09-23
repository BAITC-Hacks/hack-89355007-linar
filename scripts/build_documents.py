import json
from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
DATA = json.loads((Path(__file__).parent / 'document-data.json').read_text(encoding='utf-8'))

for period in ('before', 'after'):
    directory = ROOT / f'company_{period}'
    directory.mkdir(exist_ok=True)
    for filename, spec in DATA['files'].items():
        doc = Document()
        sec = doc.sections[0]
        sec.top_margin = sec.bottom_margin = Cm(2)
        sec.left_margin = sec.right_margin = Cm(2.3)
        normal = doc.styles['Normal']
        normal.font.name = 'Arial'
        normal.font.size = Pt(10.5)
        normal.paragraph_format.space_after = Pt(5)
        for style in ('Title', 'Heading 1'):
            doc.styles[style].font.name = 'Arial'
            doc.styles[style].font.color.rgb = __import__('docx').shared.RGBColor(0, 0, 0)
        for border in doc.styles['Title']._element.findall('.//' + qn('w:pBdr')):
            border.getparent().remove(border)
        doc.add_paragraph(filename.removesuffix('.docx').replace('Должностные инструкции', 'Должностные инструкции сотрудников').replace('Приказ о реорганизации', 'Приказ о реорганизации учебной компании'), style='Title')
        doc.add_paragraph(f'Учебная организация  |  {"до" if period == "before" else "после"} реорганизации  |  синтетический документ')
        doc.add_paragraph(f'Подразделение: {spec["unit"]}')
        doc.add_paragraph('Этот документ фиксирует функции и ответственность учебной организации для сравнения комплектов документов. Все лица и правила вымышлены.')
        doc.add_heading('Функции', level=1)
        for line in spec[period]:
            doc.add_paragraph(line)
        doc.add_heading('Полномочия', level=1)
        doc.add_paragraph(spec['authority'])
        doc.add_heading('Ответственность', level=1)
        for line in spec['responsibility'].split('\n'):
            doc.add_paragraph(line)
        doc.save(directory / filename)
        print(directory / filename)